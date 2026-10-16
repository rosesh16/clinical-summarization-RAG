from __future__ import annotations

from pathlib import Path

from .extraction_models import ExtractedDocument, ExtractedPage
from .pdf_extractor import extract_pdf
from .table_extractor import tables_from_text
from ..utils.exceptions import ExtractionError, UnsupportedDocumentError


def extract_document(file_path: str | Path) -> ExtractedDocument:
    """Extract PDF, DOCX, or TXT while retaining page/source boundaries."""
    path = Path(file_path)
    if not path.exists() or not path.is_file():
        raise ExtractionError(f"File not found: {path}")
    if path.stat().st_size == 0:
        raise ExtractionError("The uploaded file is empty.")
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".txt":
        text = path.read_text(encoding="utf-8", errors="replace")
        page = ExtractedPage(1, text, "text", tables=tables_from_text(1, text), confidence=1.0)
        return ExtractedDocument(path, path.name, [page])
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ExtractionError("DOCX support requires python-docx.") from exc
        doc = Document(path)
        paragraphs = [p.text for p in doc.paragraphs if p.text]
        table_lines = [" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
        text = "\n".join(paragraphs + table_lines)
        page = ExtractedPage(1, text, "text", tables=tables_from_text(1, text), confidence=1.0)
        return ExtractedDocument(path, path.name, [page])
    raise UnsupportedDocumentError(f"Unsupported file type '{suffix}'. Upload PDF, DOCX, or TXT.")
